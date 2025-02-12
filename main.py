import os
import requests
import streamlit as st
import PyPDF2
import pptx
from dotenv import load_dotenv


load_dotenv()
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")

def gemini_chat(prompt):
    url = "https://generativelanguage.googleapis.com/v1beta/models/gemini-pro:generateContent"
    headers = {"Content-Type": "application/json"}
    params = {"key": GEMINI_API_KEY}
    data = {"contents": [{"parts": [{"text": prompt}]}]}  # Correct request format
    
    response = requests.post(url, headers=headers, params=params, json=data)
    
    if response.status_code == 200:
        try:
            return response.json()["candidates"][0]["content"]["parts"][0]["text"]
        except (KeyError, IndexError):
            return "Invalid response format from Gemini API."
    else:
        return f"Error: {response.status_code} - {response.text}"

def extract_text_from_pdf(pdf_file):
    pdf_reader = PyPDF2.PdfReader(pdf_file)
    text = ""
    for page in pdf_reader.pages:
        text += page.extract_text() + "\n"
    return text[:7000]  # Limit text length to fit API constraints

def extract_text_from_pptx(pptx_file):
    prs = pptx.Presentation(pptx_file)
    text = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"
    return text[:7000]  # Limit text length to fit API constraints

def main():
    st.title("Chatbot")
    st.write("Ask me anything or upload a PDF/PPTX to summarize!")

    if "messages" not in st.session_state:
        st.session_state.messages = []

    for message in st.session_state.messages:
        with st.chat_message(message["role"]):
            st.markdown(message["content"])

    uploaded_file = st.file_uploader("Upload a PDF or PPTX file", type=["pdf", "pptx"])
    if uploaded_file:
        file_type = uploaded_file.type
        if file_type == "application/pdf":
            extracted_text = extract_text_from_pdf(uploaded_file)
        elif file_type == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
            extracted_text = extract_text_from_pptx(uploaded_file)
        else:
            extracted_text = "Unsupported file format."
        
        response = gemini_chat(f"Summarize this document: {extracted_text}")
        st.session_state.messages.append({"role": "assistant", "content": response})
        with st.chat_message("assistant"):
            st.markdown(response)

    prompt = st.chat_input("Type your message...")
    if prompt:
        st.session_state.messages.append({"role": "user", "content": prompt})
        with st.chat_message("user"):
            st.markdown(prompt)
        
        response = gemini_chat(prompt)
        st.session_state.messages.append({"role": "assistant", "content": response})
        with st.chat_message("assistant"):
            st.markdown(response)

if __name__ == "__main__":
    main()
